"""Uploaded CSV / XLSX files: content checks (not just the extension) and parsing.

Only `.csv` and `.xlsx` are accepted: never `.xls` or macro formats. An xlsx is a zip, so it is
checked against zip bombs (entry count, total uncompressed size) before openpyxl opens it.
"""

from __future__ import annotations

import csv
import io
import re
import unicodedata
import zipfile
from dataclasses import dataclass

import pandas as pd

from app.flows.ports import SourceError
from app.flows.sources.frame import normalize_frame

CSV_TYPE = "text/csv"
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_ZIP_ENTRIES = 1000
ZIP_BOMB_FACTOR = 10
MAX_COLUMNS = 500
_SNIFF_BYTES = 16 * 1024
MAX_HEADER_BYTES = 1024 * 1024
MAX_FIELD_CHARS = 1024 * 1024
MAX_TOTAL_CHARS = 50_000_000


class UploadRefused(SourceError):
    """The file is not acceptable (message safe to show)."""


@dataclass(frozen=True)
class ParsedFile:
    media_type: str
    sheet: str | None
    frame: pd.DataFrame


def sanitize_name(raw: str | None) -> str:
    """The stored name: no path, no control characters, at most 120 characters."""
    name = (raw or "").replace("\\", "/").rsplit("/", 1)[-1]
    name = unicodedata.normalize("NFKC", name)
    name = "".join(c for c in name if c.isprintable() and c not in '<>:"|?*')
    name = re.sub(r"\s+", " ", name).strip(" .")
    if len(name) > 120:
        stem, dot, ext = name.rpartition(".")
        name = (stem[: 115 - len(ext)] + dot + ext) if dot else name[:120]
    return name or "upload"


def file_kind(name: str) -> str:
    lower = name.lower()
    if lower.endswith(".csv"):
        return "csv"
    if lower.endswith(".xlsx"):
        return "xlsx"
    raise UploadRefused("only .csv and .xlsx files are accepted")


def check_zip(data: bytes, cap: int) -> zipfile.ZipFile:
    if not data.startswith(b"PK\x03\x04") or not zipfile.is_zipfile(io.BytesIO(data)):
        raise UploadRefused("the file is not a valid .xlsx (its content is not an Excel workbook)")
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
        infos = archive.infolist()
    except (zipfile.BadZipFile, ValueError, OSError):
        raise UploadRefused("the file is not a valid .xlsx") from None
    if len(infos) > MAX_ZIP_ENTRIES:
        raise UploadRefused("the workbook has too many internal files")
    if sum(i.file_size for i in infos) > ZIP_BOMB_FACTOR * cap:
        raise UploadRefused("the workbook is too large once unpacked")
    names = {i.filename for i in infos}
    if "xl/workbook.xml" not in names or "[Content_Types].xml" not in names:
        raise UploadRefused("the file is not an Excel workbook")
    if any("vbaproject" in n.lower() for n in names):
        raise UploadRefused("workbooks with macros are not accepted")
    if any(n.startswith("/") or ".." in n.split("/") for n in names):
        raise UploadRefused("the workbook contains unsafe paths")
    return archive


def _decode_csv(data: bytes) -> str:
    if b"\x00" in data[: _SNIFF_BYTES * 4] or data.startswith((b"PK\x03\x04", b"\xd0\xcf\x11\xe0")):
        raise UploadRefused("the file is not a text CSV")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("latin-1")


def sniff_delimiter(text: str) -> str:
    sample = text[:_SNIFF_BYTES]
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t").delimiter
    except csv.Error:
        first = sample.splitlines()[0] if sample.strip() else ""
        return max(",;\t", key=first.count)


def _clean_columns(df: pd.DataFrame) -> pd.DataFrame:
    seen: dict[str, int] = {}
    names = []
    for i, raw in enumerate(df.columns, start=1):
        name = str(raw).strip()[:200]
        if not name or name.startswith("Unnamed: "):
            name = f"column_{i}"
        count = seen.get(name, 0)
        seen[name] = count + 1
        names.append(name if count == 0 else f"{name}_{count + 1}")
    df.columns = pd.Index(names)
    return df


def _check_csv_header(text: str, sep: str) -> None:
    """Bound the work before pandas sees the file: the header line is read first (at most 1 MB)
    and its fields are counted; a wide header would otherwise cost quadratic time."""
    window = text[: MAX_HEADER_BYTES + 1]
    end = window.find("\n")
    line = window if end == -1 else window[:end]
    if len(line) > MAX_HEADER_BYTES:
        raise UploadRefused("the first line of the file is too long")
    try:
        header = next(csv.reader([line], delimiter=sep), [])
    except csv.Error:
        raise UploadRefused("the file could not be read as a table") from None
    if len(header) > MAX_COLUMNS:
        raise UploadRefused(f"the file has more than {MAX_COLUMNS} columns")
    if any(len(h) > MAX_FIELD_CHARS for h in header):
        raise UploadRefused("a field is larger than 1 MB")


def _check_field_sizes(df: pd.DataFrame) -> None:
    for col in df.columns:
        if pd.api.types.is_numeric_dtype(df[col]):
            continue
        longest = df[col].map(lambda v: len(v) if isinstance(v, str) else 0).max()
        if longest > MAX_FIELD_CHARS:
            raise UploadRefused("a field is larger than 1 MB")


def _read_xlsx(data: bytes, sheet: str | None, max_rows: int) -> tuple[pd.DataFrame, str]:
    """Stream one sheet (read-only mode); stop past max_rows rows, MAX_COLUMNS columns or
    MAX_TOTAL_CHARS characters, so a small file cannot expand into a huge frame."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
    try:
        names = list(wb.sheetnames)
        used = sheet if sheet is not None else names[0]
        if used not in names:
            raise UploadRefused(f"sheet {used!r} not found (sheets: {names[:20]})")
        rows: list[tuple] = []
        total = 0
        for row in wb[used].iter_rows(values_only=True, max_row=max_rows + 2):
            if len(row) > MAX_COLUMNS:
                while len(row) > MAX_COLUMNS and row[-1] is None:
                    row = row[:-1]
                if len(row) > MAX_COLUMNS:
                    raise UploadRefused(f"the file has more than {MAX_COLUMNS} columns")
            total += sum(len(v) if isinstance(v, str) else 8 for v in row if v is not None)
            if total > MAX_TOTAL_CHARS:
                raise UploadRefused("the sheet holds too much text")
            rows.append(row)
    finally:
        wb.close()
    rows = rows[: max_rows + 2]
    while rows and all(v is None for v in rows[-1]):
        rows.pop()
    if not rows:
        raise UploadRefused("the file is empty")
    width = max(len(r) for r in rows)
    head = [("" if v is None else v) for v in rows[0]] + [""] * (width - len(rows[0]))
    body = [list(r) + [None] * (width - len(r)) for r in rows[1 : max_rows + 2]]
    return pd.DataFrame(body, columns=head), used


def _cap(df: pd.DataFrame, max_rows: int) -> pd.DataFrame:
    if len(df) > max_rows:
        df = df.iloc[:max_rows].copy()
        df.attrs["capped"] = True
    return df


def parse_file(
    name: str, data: bytes, *, cap_bytes: int, max_rows: int, sheet: str | None = None
) -> ParsedFile:
    """Parse an upload (or a stored file). Raises UploadRefused with a message fit for the user."""
    kind = file_kind(name)
    try:
        if kind == "csv":
            text = _decode_csv(data)
            if not text.strip():
                raise UploadRefused("the file is empty")
            sep = sniff_delimiter(text)
            _check_csv_header(text, sep)
            df = pd.read_csv(
                io.StringIO(text),
                sep=sep,
                nrows=max_rows + 1,
                skip_blank_lines=True,
            )
            media, used_sheet = CSV_TYPE, None
        else:
            archive = check_zip(data, cap_bytes)
            archive.close()
            df, used_sheet = _read_xlsx(data, sheet, max_rows)
            media = XLSX_TYPE
    except UploadRefused:
        raise
    except (pd.errors.ParserError, pd.errors.EmptyDataError, UnicodeError):
        raise UploadRefused("the file could not be read as a table") from None
    except Exception:  # noqa: BLE001 - parser messages may quote cell data
        raise UploadRefused("the file could not be read as a table") from None
    if df.shape[1] == 0:
        raise UploadRefused("the file has no columns")
    if df.shape[1] > MAX_COLUMNS:
        raise UploadRefused(f"the file has more than {MAX_COLUMNS} columns")
    _check_field_sizes(df)
    df = normalize_frame(_cap(_clean_columns(df), max_rows))
    return ParsedFile(media, used_sheet, df)
